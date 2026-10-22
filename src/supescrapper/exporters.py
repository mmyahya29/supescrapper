from __future__ import annotations

import csv
import json
import re
import sqlite3
from pathlib import Path

from .exceptions import ExportError

_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def _resolve_path(dest: str | Path) -> Path:
    p = Path(dest)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _safe_table_name(table: str) -> str:
    if not _IDENT_RE.match(table):
        raise ExportError(
            f"Invalid table name {table!r}: must match {_IDENT_RE.pattern}"
        )
    return table


def _quote_ident(name: str) -> str:
    """Quote a column identifier for SQLite, escaping embedded quotes."""
    return '"' + str(name).replace('"', '""') + '"'


def _rows(results) -> list[dict]:
    """Normalize ScrapedResultSet / SchemaResultSet / item lists into dicts."""
    from .models import ScrapedResultSet

    if isinstance(results, ScrapedResultSet):
        return results.to_list()
    rows: list[dict] = []
    for r in results:
        if isinstance(r, dict):
            rows.append(r)
        elif hasattr(r, "to_dict"):
            rows.append(r.to_dict())
        else:
            rows.append({"value": r})
    return rows


def export_json(results, dest: str | Path, *, indent: int | None = 2) -> Path:
    p = _resolve_path(dest)
    data = _rows(results)
    try:
        p.write_text(json.dumps(data, indent=indent, ensure_ascii=False), encoding="utf-8")
    except Exception as exc:
        raise ExportError(f"Failed to write JSON: {exc}") from exc
    return p


def export_csv(results, dest: str | Path, *, delimiter: str = ",") -> Path:
    p = _resolve_path(dest)
    rows = _rows(results)
    if not rows:
        p.write_text("", encoding="utf-8")
        return p

    all_keys: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for k in row:
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    try:
        with p.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=all_keys, delimiter=delimiter, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
    except Exception as exc:
        raise ExportError(f"Failed to write CSV: {exc}") from exc
    return p


def export_sqlite(results, dest: str | Path, *, table: str = "results", unique_keys: list[str] | None = None) -> Path:
    p = _resolve_path(dest)
    table = _safe_table_name(table)
    rows = _rows(results)
    if not rows:
        return p

    all_keys: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for k in row:
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    quoted_table = _quote_ident(table)
    cols_sql = ", ".join(f"{_quote_ident(k)} TEXT" for k in all_keys)
    col_names_sql = ", ".join(_quote_ident(k) for k in all_keys)
    placeholders = ", ".join("?" for _ in all_keys)

    try:
        with sqlite3.connect(str(p)) as conn:
            cur = conn.cursor()
            cur.execute(f"CREATE TABLE IF NOT EXISTS {quoted_table} ({cols_sql})")

            insert_verb = "INSERT"
            if unique_keys:
                valid_unique = [k for k in unique_keys if k in all_keys]
                if valid_unique:
                    cols_for_idx = ", ".join(_quote_ident(k) for k in valid_unique)
                    idx_name = _quote_ident(f"uq_{table}_{'_'.join(valid_unique)}")
                    cur.execute(
                        f"CREATE UNIQUE INDEX IF NOT EXISTS {idx_name} "
                        f"ON {quoted_table} ({cols_for_idx})"
                    )
                    insert_verb = "INSERT OR REPLACE"

            insert_sql = f"{insert_verb} INTO {quoted_table} ({col_names_sql}) VALUES ({placeholders})"
            for row in rows:
                values = [row.get(k, "") for k in all_keys]
                cur.execute(insert_sql, values)
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError(f"Failed to write SQLite: {exc}") from exc
    return p


def export_xlsx(results, dest: str | Path) -> Path:
    """Export to Excel (.xlsx) format."""
    p = _resolve_path(dest)
    rows = _rows(results)
    if not rows:
        p.write_bytes(b"")
        return p

    try:
        import openpyxl
    except ImportError as exc:
        raise ExportError("openpyxl is required for XLSX export. Install with: pip install openpyxl") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()

    all_keys: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for k in row:
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    ws.append(all_keys)
    for row in rows:
        ws.append([row.get(k, "") for k in all_keys])

    try:
        wb.save(str(p))
    except Exception as exc:
        raise ExportError(f"Failed to write XLSX: {exc}") from exc
    return p


def export_records(results, fmt: str, dest: str | Path) -> Path:
    """Dispatch export based on format."""
    fmt = fmt.lower()
    if fmt in ("json", "jsonl"):
        return export_json(results, dest)
    if fmt == "csv":
        return export_csv(results, dest)
    if fmt in ("sqlite", "sql"):
        return export_sqlite(results, dest)
    if fmt in ("xlsx", "xls", "excel"):
        return export_xlsx(results, dest)
    raise ExportError(f"Unsupported export format: {fmt}")
